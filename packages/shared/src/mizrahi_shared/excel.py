"""
Excel report generation utilities.

This module provides utilities for creating styled Excel reports with
multiple sheets, consistent formatting, and review columns.
"""

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# Style constants
HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
PASS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Review columns (Hebrew)
REVIEW_COLUMNS = [
    ("האם תקין?", "Is OK?"),
    ("שם הבודק", "Checker Name"),
]


class ExcelGenerator:
    """
    Excel report generator with consistent styling.

    Usage:
        generator = ExcelGenerator()
        generator.add_sheet("סיכום", summary_data, columns=["שדה", "ערך"])
        generator.add_sheet("ממצאים", findings, include_review_columns=True)
        generator.save("report.xlsx")
    """

    def __init__(self, include_review_columns: bool = True):
        """
        Initialize Excel generator.

        Args:
            include_review_columns: Whether to add review columns by default.
        """
        self.workbook = Workbook()
        self.include_review_columns = include_review_columns
        self._sheet_count = 0

        # Remove default sheet
        if "Sheet" in self.workbook.sheetnames:
            del self.workbook["Sheet"]

    def add_sheet(
        self,
        name: str,
        data: List[Dict[str, Any]],
        columns: Optional[List[str]] = None,
        include_review_columns: Optional[bool] = None,
    ) -> Worksheet:
        """
        Add a sheet with data.

        Args:
            name: Sheet name.
            data: List of dictionaries (rows).
            columns: Column headers. If not provided, uses keys from first data row.
            include_review_columns: Override default review columns setting.

        Returns:
            The created worksheet.
        """
        # Truncate name if too long (Excel limit is 31 characters)
        if len(name) > 31:
            name = name[:28] + "..."

        ws = self.workbook.create_sheet(title=name)
        self._sheet_count += 1

        if not data:
            ws.cell(row=1, column=1, value="אין נתונים")
            return ws

        # Determine columns
        if columns is None:
            columns = list(data[0].keys())

        # Add review columns if needed
        add_review = include_review_columns if include_review_columns is not None else self.include_review_columns
        if add_review:
            columns = columns + [rc[0] for rc in REVIEW_COLUMNS]

        # Write headers
        for col_idx, header in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER

        # Write data
        for row_idx, row_data in enumerate(data, start=2):
            for col_idx, header in enumerate(columns, start=1):
                # Skip review columns in data (they're empty for user input)
                if header in [rc[0] for rc in REVIEW_COLUMNS]:
                    cell = ws.cell(row=row_idx, column=col_idx, value="")
                else:
                    value = row_data.get(header, "")
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)

                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.border = THIN_BORDER

        # Auto-fit columns
        self._auto_fit_columns(ws)

        # Set RTL direction
        ws.sheet_view.rightToLeft = True

        return ws

    def add_summary_sheet(
        self,
        hook_name: str,
        manager_name: str,
        check_results: List[Dict[str, Any]],
        timestamp: Optional[datetime] = None,
    ) -> Worksheet:
        """
        Add a summary sheet with hook execution results.

        Args:
            hook_name: Name of the hook.
            manager_name: Manager name.
            check_results: List of check results.
            timestamp: Execution timestamp.

        Returns:
            The created worksheet.
        """
        ws = self.workbook.create_sheet(title="סיכום", index=0)

        # Title
        ws.merge_cells("A1:D1")
        title_cell = ws.cell(row=1, column=1, value=hook_name)
        title_cell.font = Font(bold=True, size=16)
        title_cell.alignment = Alignment(horizontal="center")

        # Info section
        info_data = [
            ("מנהל הקרן:", manager_name),
            ("תאריך הרצה:", (timestamp or datetime.now()).strftime("%Y-%m-%d %H:%M:%S")),
            ("סה״כ בדיקות:", len(check_results)),
            ("בדיקות שעברו:", sum(1 for c in check_results if c.get("status") == "pass")),
            ("בדיקות שנכשלו:", sum(1 for c in check_results if c.get("status") == "fail")),
            ("אזהרות:", sum(1 for c in check_results if c.get("status") == "warning")),
        ]

        for row_idx, (label, value) in enumerate(info_data, start=3):
            ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row_idx, column=2, value=value)

        # Check results table
        start_row = len(info_data) + 5
        headers = ["בדיקה", "סטטוס", "ממצאים", "הערות"]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        for row_idx, check in enumerate(check_results, start=start_row + 1):
            status = check.get("status", "")
            status_he = {"pass": "עבר", "fail": "נכשל", "warning": "אזהרה", "skipped": "דולג"}.get(status, status)

            ws.cell(row=row_idx, column=1, value=check.get("check_name_he", "")).border = THIN_BORDER
            status_cell = ws.cell(row=row_idx, column=2, value=status_he)
            status_cell.border = THIN_BORDER

            # Color code status
            if status == "pass":
                status_cell.fill = PASS_FILL
            elif status == "fail":
                status_cell.fill = FAIL_FILL
            elif status == "warning":
                status_cell.fill = WARNING_FILL

            ws.cell(row=row_idx, column=3, value=check.get("findings_count", 0)).border = THIN_BORDER
            ws.cell(row=row_idx, column=4, value=check.get("message", "")).border = THIN_BORDER

        self._auto_fit_columns(ws)
        ws.sheet_view.rightToLeft = True

        return ws

    def add_check_status_sheet(
        self,
        check_results: List[Dict[str, Any]],
    ) -> Worksheet:
        """
        Add a sheet with check status details.

        Args:
            check_results: List of check results.

        Returns:
            The created worksheet.
        """
        data = [
            {
                "מזהה בדיקה": c.get("check_id", ""),
                "שם הבדיקה": c.get("check_name_he", ""),
                "סטטוס": {"pass": "עבר", "fail": "נכשל", "warning": "אזהרה", "skipped": "דולג"}.get(
                    c.get("status", ""), c.get("status", "")
                ),
                "מספר ממצאים": c.get("findings_count", 0),
                "משך (מ״ש)": c.get("duration_ms", 0),
                "הערות": c.get("message", ""),
            }
            for c in check_results
        ]

        return self.add_sheet("סטטוס בדיקות", data, include_review_columns=False)

    def _auto_fit_columns(self, ws: Worksheet, min_width: int = 10, max_width: int = 50):
        """Auto-fit column widths based on content."""
        for column_cells in ws.columns:
            max_length = 0
            column = column_cells[0].column_letter

            for cell in column_cells:
                try:
                    cell_length = len(str(cell.value or ""))
                    if cell_length > max_length:
                        max_length = cell_length
                except:
                    pass

            adjusted_width = max(min_width, min(max_length + 2, max_width))
            ws.column_dimensions[column].width = adjusted_width

    def save(self, filepath: Union[str, Path]) -> Path:
        """
        Save the workbook to a file.

        Args:
            filepath: Path to save the file.

        Returns:
            Path to the saved file.
        """
        filepath = Path(filepath)
        filepath.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(filepath)
        logger.info(f"Excel report saved: {filepath}")
        return filepath


async def generate_report(
    output_path: Path,
    hook_name: str,
    manager_name: str,
    check_results: List[Dict[str, Any]],
    findings_by_check: Dict[str, List[Dict[str, Any]]],
    include_review_columns: bool = True,
) -> Path:
    """
    Generate a complete Excel report.

    Args:
        output_path: Path to save the report.
        hook_name: Name of the hook.
        manager_name: Manager name.
        check_results: List of check results.
        findings_by_check: Dictionary mapping check IDs to their findings.
        include_review_columns: Whether to include review columns.

    Returns:
        Path to the generated report.
    """
    generator = ExcelGenerator(include_review_columns=include_review_columns)

    # Add summary sheet
    generator.add_summary_sheet(
        hook_name=hook_name,
        manager_name=manager_name,
        check_results=check_results,
    )

    # Add check status sheet
    generator.add_check_status_sheet(check_results)

    # Add findings sheets
    for check in check_results:
        check_id = check.get("check_id", "")
        findings = findings_by_check.get(check_id, [])

        if findings:
            sheet_name = check.get("check_name_he", check_id)
            generator.add_sheet(sheet_name, findings)

    return generator.save(output_path)
