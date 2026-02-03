#!/usr/bin/env python3
"""
Compare new output with legacy output to ensure exact match.

Usage:
    python compare_output.py --new-file output.xlsx --legacy-file expected.xlsx
    python compare_output.py --new-dir ./output --legacy-dir ./test_data/hook2_legacy_complete
"""

import sys
import argparse
from pathlib import Path
from typing import Dict, List, Tuple
import json

try:
    import pandas as pd
    from openpyxl import load_workbook
except ImportError:
    print("Missing required packages. Install with:")
    print("pip install pandas openpyxl")
    sys.exit(1)


def compare_excel_files(new_path: Path, legacy_path: Path) -> Dict:
    """Compare two Excel files and return differences"""
    result = {
        "match": True,
        "new_file": str(new_path),
        "legacy_file": str(legacy_path),
        "differences": [],
        "summary": {},
    }

    try:
        # Load workbooks
        new_wb = load_workbook(new_path)
        legacy_wb = load_workbook(legacy_path)

        new_sheets = set(new_wb.sheetnames)
        legacy_sheets = set(legacy_wb.sheetnames)

        # Check sheet names
        only_in_new = new_sheets - legacy_sheets
        only_in_legacy = legacy_sheets - new_sheets
        common_sheets = new_sheets & legacy_sheets

        if only_in_new:
            result["differences"].append(f"Sheets only in new: {only_in_new}")
            result["match"] = False

        if only_in_legacy:
            result["differences"].append(f"Sheets only in legacy: {only_in_legacy}")
            result["match"] = False

        result["summary"]["new_sheets"] = len(new_sheets)
        result["summary"]["legacy_sheets"] = len(legacy_sheets)
        result["summary"]["common_sheets"] = len(common_sheets)

        # Compare common sheets
        for sheet_name in common_sheets:
            new_df = pd.read_excel(new_path, sheet_name=sheet_name)
            legacy_df = pd.read_excel(legacy_path, sheet_name=sheet_name)

            # Check row counts
            if len(new_df) != len(legacy_df):
                result["differences"].append(
                    f"Sheet '{sheet_name}': Row count differs - new: {len(new_df)}, legacy: {len(legacy_df)}"
                )
                result["match"] = False

            # Check column counts
            if len(new_df.columns) != len(legacy_df.columns):
                result["differences"].append(
                    f"Sheet '{sheet_name}': Column count differs - new: {len(new_df.columns)}, legacy: {len(legacy_df.columns)}"
                )
                result["match"] = False

            # Check column names
            new_cols = set(new_df.columns)
            legacy_cols = set(legacy_df.columns)
            col_diff = new_cols.symmetric_difference(legacy_cols)
            if col_diff:
                result["differences"].append(
                    f"Sheet '{sheet_name}': Column name differences: {col_diff}"
                )
                result["match"] = False

            # Check data content (for sheets with same structure)
            if len(new_df) == len(legacy_df) and len(new_df.columns) == len(legacy_df.columns):
                common_cols = list(new_cols & legacy_cols)
                if common_cols:
                    # Compare data values
                    for col in common_cols:
                        new_col = new_df[col].fillna("").astype(str)
                        legacy_col = legacy_df[col].fillna("").astype(str)

                        if not new_col.equals(legacy_col):
                            # Count differences
                            diff_mask = new_col != legacy_col
                            diff_count = diff_mask.sum()
                            if diff_count > 0:
                                result["differences"].append(
                                    f"Sheet '{sheet_name}', Column '{col}': {diff_count} value(s) differ"
                                )
                                # Show first few differences
                                diff_indices = diff_mask[diff_mask].index[:3]
                                for idx in diff_indices:
                                    result["differences"].append(
                                        f"  Row {idx}: new='{new_col[idx]}' vs legacy='{legacy_col[idx]}'"
                                    )
                                result["match"] = False

    except Exception as e:
        result["match"] = False
        result["differences"].append(f"Error comparing files: {str(e)}")

    return result


def compare_json_files(new_path: Path, legacy_path: Path) -> Dict:
    """Compare two JSON files"""
    result = {
        "match": True,
        "new_file": str(new_path),
        "legacy_file": str(legacy_path),
        "differences": [],
    }

    try:
        with open(new_path, encoding="utf-8") as f:
            new_data = json.load(f)
        with open(legacy_path, encoding="utf-8") as f:
            legacy_data = json.load(f)

        if new_data != legacy_data:
            result["match"] = False
            result["differences"].append(f"JSON content differs")
            result["new_data"] = new_data
            result["legacy_data"] = legacy_data

    except Exception as e:
        result["match"] = False
        result["differences"].append(f"Error comparing JSON: {str(e)}")

    return result


def compare_directories(new_dir: Path, legacy_dir: Path, manager: str = None) -> Dict:
    """Compare all output files in two directories"""
    results = {
        "overall_match": True,
        "file_comparisons": [],
        "summary": {
            "total_files": 0,
            "matching": 0,
            "different": 0,
            "missing_in_new": 0,
            "missing_in_legacy": 0,
        },
    }

    # Find Excel and JSON files
    if manager:
        new_xlsx = list(new_dir.glob(f"**/{manager}*_report.xlsx"))
        legacy_xlsx = list(legacy_dir.glob(f"**/{manager}*_report.xlsx"))
        new_json = list(new_dir.glob(f"**/{manager}*_email.json"))
        legacy_json = list(legacy_dir.glob(f"**/{manager}*_email.json"))
    else:
        new_xlsx = list(new_dir.glob("**/*_report.xlsx"))
        legacy_xlsx = list(legacy_dir.glob("**/*_report.xlsx"))
        new_json = list(new_dir.glob("**/*_email.json"))
        legacy_json = list(legacy_dir.glob("**/*_email.json"))

    # Match files by name
    new_files = {f.name: f for f in new_xlsx + new_json}
    legacy_files = {f.name: f for f in legacy_xlsx + legacy_json}

    all_names = set(new_files.keys()) | set(legacy_files.keys())
    results["summary"]["total_files"] = len(all_names)

    for name in all_names:
        if name not in new_files:
            results["file_comparisons"].append(
                {"file": name, "status": "missing_in_new", "legacy_path": str(legacy_files[name])}
            )
            results["summary"]["missing_in_new"] += 1
            results["overall_match"] = False
        elif name not in legacy_files:
            results["file_comparisons"].append(
                {"file": name, "status": "missing_in_legacy", "new_path": str(new_files[name])}
            )
            results["summary"]["missing_in_legacy"] += 1
            # Not necessarily a failure - could be new functionality
        else:
            # Compare files
            new_path = new_files[name]
            legacy_path = legacy_files[name]

            if name.endswith(".xlsx"):
                comparison = compare_excel_files(new_path, legacy_path)
            elif name.endswith(".json"):
                comparison = compare_json_files(new_path, legacy_path)
            else:
                comparison = {"match": True, "skipped": True}

            comparison["file"] = name
            comparison["status"] = "match" if comparison["match"] else "different"
            results["file_comparisons"].append(comparison)

            if comparison["match"]:
                results["summary"]["matching"] += 1
            else:
                results["summary"]["different"] += 1
                results["overall_match"] = False

    return results


def print_report(results: Dict):
    """Print comparison report"""
    print("\n" + "=" * 60)
    print("COMPARISON REPORT")
    print("=" * 60)

    summary = results.get("summary", {})
    print(f"\nTotal files: {summary.get('total_files', 0)}")
    print(f"Matching: {summary.get('matching', 0)}")
    print(f"Different: {summary.get('different', 0)}")
    print(f"Missing in new: {summary.get('missing_in_new', 0)}")
    print(f"Missing in legacy: {summary.get('missing_in_legacy', 0)}")

    overall = "PASS" if results.get("overall_match", False) else "FAIL"
    print(f"\nOverall result: {overall}")
    print("=" * 60)

    # Print details for non-matching files
    for comp in results.get("file_comparisons", []):
        if comp.get("status") != "match":
            print(f"\n--- {comp.get('file', 'unknown')} ---")
            print(f"Status: {comp.get('status', 'unknown')}")
            for diff in comp.get("differences", []):
                print(f"  - {diff}")

    print("\n" + "=" * 60)


def main():
    parser = argparse.ArgumentParser(description="Compare output with legacy")
    parser.add_argument("--new-file", type=Path, help="New output file to compare")
    parser.add_argument("--legacy-file", type=Path, help="Legacy expected file")
    parser.add_argument("--new-dir", type=Path, help="New output directory")
    parser.add_argument("--legacy-dir", type=Path, help="Legacy output directory")
    parser.add_argument("--manager", help="Manager name to compare (optional)")
    parser.add_argument("--output-json", type=Path, help="Save results to JSON")

    args = parser.parse_args()

    if args.new_file and args.legacy_file:
        # Compare single files
        if args.new_file.suffix == ".xlsx":
            results = compare_excel_files(args.new_file, args.legacy_file)
        elif args.new_file.suffix == ".json":
            results = compare_json_files(args.new_file, args.legacy_file)
        else:
            print(f"Unsupported file type: {args.new_file.suffix}")
            sys.exit(1)

        # Wrap in structure for print_report
        results = {
            "overall_match": results.get("match", False),
            "file_comparisons": [results],
            "summary": {
                "total_files": 1,
                "matching": 1 if results.get("match") else 0,
                "different": 0 if results.get("match") else 1,
            },
        }

    elif args.new_dir and args.legacy_dir:
        # Compare directories
        results = compare_directories(args.new_dir, args.legacy_dir, args.manager)
    else:
        print("Provide either --new-file/--legacy-file or --new-dir/--legacy-dir")
        parser.print_help()
        sys.exit(1)

    print_report(results)

    if args.output_json:
        with open(args.output_json, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=2, default=str)
        print(f"\nResults saved to: {args.output_json}")

    # Exit with error code if not matching
    sys.exit(0 if results.get("overall_match", False) else 1)


if __name__ == "__main__":
    main()
